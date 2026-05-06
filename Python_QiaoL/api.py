from pptx import Presentation

# 读取PPTX文件
def read_pptx(file_path):
    prs = Presentation(file_path)
    text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text.append(shape.text)
    return "\n".join(text)

# 将文字写入TXT文件
def write_to_txt(text, output_file_path):
    with open(output_file_path, "w", encoding="utf-8") as file:
        file.write(text)

# 主函数
def main():
    pptx_file_path = "D:\\1.Introduction.pptx"
    output_file_path = "D:\\output.txt"
    text = read_pptx(pptx_file_path)
    write_to_txt(text, output_file_path)
    print(f"Text extracted from {pptx_file_path} and saved to {output_file_path}")

if __name__ == "__main__":
    main()
